from datetime import datetime
from unittest.mock import patch

from openpyxl import load_workbook

import apartments
import main


def catalog_item(
    item_id=1,
    *,
    description="Продаю свою квартиру. Собственник.",
    seller_label=None,
    closed="3 завершённых объявления",
    phone=True,
):
    item = {
        "id": item_id,
        "type": "item",
        "title": "2-к. квартира, 54,5 м², 3/9 эт.",
        "urlPath": f"/sankt-peterburg/kvartiry/2-k._kvartira_{item_id}?context=abc",
        "description": description,
        "priceDetailed": {"value": 10_900_000},
        "coords": {"address_user": "Санкт-Петербург, Примерная ул., 1"},
        "geo": {
            "formattedAddress": "Примерная ул., 1",
            "geoReferences": [{"content": "Примерная", "after": " 500 м"}],
        },
        "sortTimeStamp": 1_790_000_000_000,
        "closedItemsText": closed,
        "contacts": {"phone": phone},
        "iva": {
            "DateInfoStep": [
                {"componentData": {"component": "date-info"}, "payload": {"relative": "2 часа назад"}}
            ],
            "BadgeBarStep": [
                {"payload": {"badges": [{"title": "Собственник"}, {"title": "Проверено в Росреестре"}]}}
            ],
            "SecondLineStep": [],
        },
    }
    if seller_label:
        item["iva"]["SecondLineStep"] = [{"payload": {"value": seller_label}}]
    return item


def classified(**kwargs):
    return apartments.classify(apartments.listing_from_item(catalog_item(**kwargs)))


def test_listing_fields_are_extracted():
    listing = apartments.listing_from_item(catalog_item(7))

    assert listing.id == 7
    assert listing.url == "https://www.avito.ru/sankt-peterburg/kvartiry/2-k._kvartira_7"
    assert listing.price == 10_900_000
    assert listing.area == 54.5
    assert listing.floor == "3/9"
    assert listing.price_per_meter == 200_000
    assert listing.address == "Санкт-Петербург, Примерная ул., 1"
    assert listing.metro == "Примерная 500 м"
    assert listing.published == "2 часа назад"
    assert listing.closed_listings == 3
    assert listing.calls_allowed is True
    assert listing.avito_badges == ("Собственник", "Проверено в Росреестре")


def test_plain_owner_listing_is_owner():
    listing = classified()

    assert listing.verdict == apartments.VERDICT_OWNER
    assert "пишет «собственник»" in listing.reasons


def test_avito_agency_label_rejects_even_with_owner_badge_and_text():
    listing = classified(
        seller_label="Агентство",
        description="Продажа напрямую от собственника.",
    )

    assert listing.verdict == apartments.VERDICT_AGENT


def test_agency_phrases_in_description_reject():
    for text in (
        "Эксклюзивные права продажи у нашего агентства.",
        "Наше агентство недвижимости работает с 2005 года.",
        "Бесплатное юридическое сопровождение сделки.",
        "Собственники с другими агентами не сотрудничают.",
        "Комиссия агентства 2%.",
    ):
        assert classified(description=text).verdict == apartments.VERDICT_AGENT, text


def test_owner_asking_agents_not_to_call_stays_owner():
    listing = classified(
        description=(
            "Прошу агентов и риэлторов не беспокоить. "
            "Агентам без реальных покупателей не звонить. "
            "От собственника без комиссии."
        )
    )

    assert listing.verdict == apartments.VERDICT_OWNER
    assert listing.asks_no_agents is True


def test_owner_refusals_are_not_agent_signals():
    for text in (
        "Эксклюзивный договор не подписываю.",
        "Комиссию не плачу.",
        "Прямая продажа, без скрытых комиссий.",
        "Юридическое сопровождение сделки продавцом.",
        "Эксклюзивное предложение в центре.",
    ):
        assert classified(description=text).verdict == apartments.VERDICT_OWNER, text


def test_owner_ready_to_pay_agent_is_flagged():
    listing = classified(
        description="Готов заплатить комиссию агенту, который приведет покупателя."
    )

    assert listing.verdict == apartments.VERDICT_OWNER
    assert listing.pays_agent is True


def test_owner_with_own_agent_is_flagged_but_kept():
    for text in (
        "Агент на сопровождение есть. Звоните.",
        "Риэлтор для сопровождения сделки есть, в услугах не нуждаюсь.",
    ):
        listing = classified(description=text)
        assert listing.verdict == apartments.VERDICT_OWNER, text
        assert listing.has_own_agent is True, text


def test_many_closed_listings_needs_manual_check():
    listing = classified(closed="150 завершённых объявлений")

    assert listing.verdict == apartments.VERDICT_CHECK


def test_new_listings_are_marked_against_seen_registry():
    old, new = (apartments.listing_from_item(catalog_item(i)) for i in (1, 2))
    seen = {"1": "2026-09-01T10:00:00"}

    updated = apartments.mark_new([old, new], seen, datetime(2026, 9, 30, 12, 0))

    assert (old.is_new, new.is_new) == (False, True)
    assert updated == {"1": "2026-09-01T10:00:00", "2": "2026-09-30T12:00:00"}


def test_workbook_splits_owners_and_agents(tmp_path):
    owner = classified(item_id=1)
    agent = classified(item_id=2, seller_label="Агентство")
    path = tmp_path / "out.xlsx"

    apartments.write_workbook(path, [agent, owner])

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Собственники", "Отсеяно — агенты"]
    owners = workbook["Собственники"]
    assert owners.max_row == 2
    link_column = [c.value for c in owners[1]].index("Ссылка") + 1
    assert owners.cell(row=2, column=link_column).hyperlink.target.endswith("_1")
    assert workbook["Отсеяно — агенты"].max_row == 2


def test_page_with_only_seen_listings_stops_incremental_walk():
    result = main.PageRequestResult(
        page=1, status_code=200, items=({"id": 1}, {"id": 2})
    )

    assert apartments.page_has_nothing_new(result, {"1": "x", "2": "x"})
    assert not apartments.page_has_nothing_new(result, {"1": "x"})


def test_full_walk_resumes_after_last_successful_page():
    state = {"complete": False, "next_page": 1}
    stopped = apartments.ScanOutcome(items=[], error="GeeTest", last_ok_page=13, reached_end=False)
    finished = apartments.ScanOutcome(items=[], error=None, last_ok_page=22, reached_end=True)

    assert apartments.next_state(state, stopped, full_walk=True) == {"complete": False, "next_page": 14}
    assert apartments.next_state(state, finished, full_walk=True) == {"complete": True, "next_page": 1}


def test_collect_listings_keeps_pages_gathered_before_a_failure():
    stats = main.ItemsPageStats(
        count=100, total_count=100, total_elements=100, main_count=100,
        items_on_page=50, items_hash=None,
    )

    def fake_run(on_page_results, start_page, stop_after):
        on_page_results((
            main.PageRequestResult(page=start_page, status_code=200, stats=stats, items=({"id": 5},)),
        ))
        raise RuntimeError("GeeTest failed 5 consecutive times")

    with patch.object(main, "run", side_effect=fake_run):
        outcome = apartments.collect_listings(start_page=1)

    assert [item["id"] for item in outcome.items] == [5]
    assert outcome.error == "GeeTest failed 5 consecutive times"
    assert outcome.last_ok_page == 1
    assert outcome.reached_end is False
