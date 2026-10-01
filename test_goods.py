from openpyxl import load_workbook

import avito_client
import goods
import profiles


def phone(item_id, price, *, title="iPhone 15, 128 ГБ, SIM + eSIM", city="sankt-peterburg",
          description="Продаю свой телефон.", seller=None, closed="5 завершённых объявлений", badges=()):
    item = {
        "id": item_id,
        "type": "item",
        "title": title,
        "urlPath": f"/{city}/telefony/iphone_15_{item_id}",
        "description": description,
        "priceDetailed": {"value": price},
        "closedItemsText": closed,
        "contacts": {"phone": True},
        "iva": {
            "BadgeBarStep": [{"payload": {"badges": [{"title": b} for b in badges]}}],
            "SecondLineStep": [{"payload": {"value": seller}}] if seller else [],
        },
    }
    return item


def test_title_gives_model_memory_and_city():
    offer = goods.offer_from_item(phone(1, 35_000, city="moskva", badges=("Только доставка", "Цена ниже рыночной")))

    assert offer.group == "iPhone 15, 128 ГБ"
    assert offer.city == "moskva"
    assert offer.delivery_only is True
    assert offer.avito_price_badge == "Цена ниже рыночной"
    assert offer.seller == "частное лицо"


def test_market_median_ignores_defects_and_extremes():
    offers = [goods.offer_from_item(phone(i, price)) for i, price in enumerate((34_000, 35_000, 36_000, 37_000))]
    offers.append(goods.offer_from_item(phone(10, 9_000)))  # extreme
    offers.append(goods.offer_from_item(phone(11, 15_000, description="Заблокирован по iCloud")))

    stats = goods.market_stats(offers)["iPhone 15, 128 ГБ"]

    assert stats.count == 5
    assert stats.used == 4
    assert stats.median == 35_500


def test_annotations_mark_cheap_defect_shop_and_other_city():
    offers = [goods.offer_from_item(phone(i, 36_000)) for i in range(4)]
    cheap = goods.offer_from_item(phone(20, 29_000))
    locked = goods.offer_from_item(phone(21, 15_000, description="Битый экран, на запчасти"))
    shop = goods.offer_from_item(phone(22, 36_000, closed="777 завершённых объявлений"))
    remote = goods.offer_from_item(phone(23, 36_000, city="moskva"))
    everything = offers + [cheap, locked, shop, remote]
    stats = goods.market_stats(everything)

    goods.annotate(everything, stats, "sankt-peterburg")

    assert "дешевле рынка" in cheap.notes
    assert locked.defect and "заблокирован/неисправен (по описанию)" in locked.notes
    assert any("похоже на магазин" in note for note in shop.notes)
    assert "другой город: moskva" in remote.notes


def test_not_broken_is_not_a_defect():
    assert not goods.offer_from_item(phone(1, 35_000, description="Не битый, не вскрывался")).defect


def test_goods_profile_alerts_only_for_home_city(tmp_path):
    base = {str(i): phone(i, 36_000) for i in range(4)}
    profile = profiles.GoodsProfile(base, "sankt-peterburg")

    assert profile.alert(phone(50, 30_000, city="moskva")) is None
    alert = profile.alert(phone(51, 30_000))
    assert alert["price"] == "30 000 ₽"
    assert alert["vs_market"] == "-17%"

    path = tmp_path / "phones.xlsx"
    profile.write_table(path, seen_before={})
    assert load_workbook(path).sheetnames == ["Рынок", "Объявления"]


def test_short_page_is_last():
    full = avito_client.Page(number=1, items=({"id": 1},), total_count=1065, items_on_page=50)
    short = avito_client.Page(number=22, items=({"id": 1},), total_count=1065, items_on_page=15)
    empty = avito_client.Page(number=23, items=(), total_count=1065, items_on_page=0)

    assert not full.is_last
    assert short.is_last
    assert empty.is_last


def test_items_url_keeps_search_params_and_page():
    url = avito_client.items_url({"categoryId": "84", "q": "iphone 15"}, 3)

    assert "categoryId=84" in url and "q=iphone+15" in url and "p=3" in url


def test_goods_alerts_only_private_people_by_default():
    base = {str(i): phone(i, 36_000) for i in range(4)}
    private = profiles.GoodsProfile(base, "sankt-peterburg")
    everyone = profiles.GoodsProfile(base, "sankt-peterburg", alert_sellers="all")

    company = phone(60, 36_000, seller="Компания")
    shop = phone(61, 36_000, closed="8884 завершённых объявления")
    person = phone(62, 36_000)

    assert private.alert(company) is None
    assert private.alert(shop) is None
    assert private.alert(person) is not None
    assert everyone.alert(company) is not None
