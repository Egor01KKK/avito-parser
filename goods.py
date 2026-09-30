"""Goods profile: group listings by model and memory, compute market prices.

Avito titles in structured categories look like "iPhone 15, 128 ГБ, SIM +
eSIM": the first part is the model, one part is the storage. A group is
"model + storage"; its market price is the median after dropping extreme
prices, which are usually broken devices, parts, scams or typos.
"""

from __future__ import annotations

import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import apartments

MEMORY_PATTERN = re.compile(r"^\d+\s*(?:ГБ|ТБ|GB|TB)$", re.IGNORECASE)
# Prices outside [LOW, HIGH] x raw median are left out of the market price.
OUTLIER_LOW = 0.5
OUTLIER_HIGH = 1.8
CHEAP_SHARE = 0.85
MIN_GROUP_SIZE = 3
# A "private" seller with this many closed listings is a shop or a reseller.
SHOP_CLOSED_LISTINGS = 100

# Locked or broken devices are sold far below market; they are marked and
# left out of the market price.
DEFECT_PATTERN = re.compile(
    r"(?<!не\s)бит(?:ый|ая|ое|ые)\b|заблок|блокировк|icloud|айклауд|аклауд|на\s+запчаст|на\s+разбор|под\s+восстановл|"
    r"не\s+включ|не\s+работает|разбит|треснут|утоплен|после\s+воды|залит|"
    r"не\s+ловит|плата\s+мертв|донор",
    re.IGNORECASE,
)


@dataclass
class Offer:
    listing: apartments.Listing
    model: str
    memory: str
    seller: str
    delivery_only: bool
    avito_price_badge: str
    city: str = ""
    defect: bool = False
    market: int | None = None
    vs_market: float | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def group(self) -> str:
        return f"{self.model}, {self.memory}" if self.memory else self.model


def offer_from_item(item: dict[str, Any]) -> Offer:
    listing = apartments.listing_from_item(item)
    parts = [part.strip() for part in listing.title.split(",") if part.strip()]
    model = parts[0] if parts else listing.title
    memory = next((part for part in parts[1:] if MEMORY_PATTERN.match(part)), "")
    badges = listing.avito_badges
    price_badge = next(
        (badge for badge in badges if "рыночн" in badge.lower()),
        "",
    )
    return Offer(
        listing=listing,
        model=model,
        memory=memory,
        seller=listing.seller_label or "частное лицо",
        delivery_only="Только доставка" in badges,
        avito_price_badge=price_badge,
        city=str(item.get("urlPath") or "").strip("/").split("/")[0],
        defect=bool(DEFECT_PATTERN.search(f"{listing.title} {listing.description}")),
    )


@dataclass(frozen=True)
class GroupStats:
    group: str
    count: int
    used: int
    low: int
    q1: int
    median: int
    q3: int
    high: int


def market_stats(offers: list[Offer]) -> dict[str, GroupStats]:
    groups: dict[str, list[int]] = {}
    for offer in offers:
        if offer.listing.price and not offer.defect:
            groups.setdefault(offer.group, []).append(offer.listing.price)
    result = {}
    for group, prices in groups.items():
        if len(prices) < MIN_GROUP_SIZE:
            continue
        raw_median = statistics.median(prices)
        kept = sorted(
            price for price in prices
            if OUTLIER_LOW * raw_median <= price <= OUTLIER_HIGH * raw_median
        )
        if len(kept) < MIN_GROUP_SIZE:
            continue
        q1, median, q3 = statistics.quantiles(kept, n=4)
        result[group] = GroupStats(
            group=group, count=len(prices), used=len(kept),
            low=kept[0], q1=round(q1), median=round(median), q3=round(q3), high=kept[-1],
        )
    return result


def annotate(
    offers: list[Offer], stats: dict[str, GroupStats], home_city: str = ""
) -> None:
    for offer in offers:
        group = stats.get(offer.group)
        offer.notes = []
        if group and offer.listing.price:
            offer.market = group.median
            offer.vs_market = offer.listing.price / group.median - 1
            if offer.defect:
                offer.notes.append("заблокирован/неисправен (по описанию)")
            elif offer.listing.price < OUTLIER_LOW * group.median:
                offer.notes.append("подозрительно дёшево: битый, запчасти или обман?")
            elif offer.listing.price <= CHEAP_SHARE * group.median:
                offer.notes.append("дешевле рынка")
        if offer.seller != "частное лицо":
            offer.notes.append(f"продавец: {offer.seller}")
        elif (offer.listing.closed_listings or 0) >= SHOP_CLOSED_LISTINGS:
            offer.notes.append(
                f"похоже на магазин/перекупа: {offer.listing.closed_listings} завершённых"
            )
        if offer.delivery_only:
            offer.notes.append("только доставка")
        if home_city and offer.city and offer.city != home_city:
            offer.notes.append(f"другой город: {offer.city}")


COLUMNS = (
    ("Новое", 8), ("Группа", 22), ("Цена, ₽", 11), ("Рынок (медиана), ₽", 13),
    ("К рынку", 9), ("Оценка Авито", 18), ("Заметки", 40), ("Продавец", 14),
    ("Завершённых объявлений", 11), ("Опубликовано", 14), ("Адрес", 30),
    ("Заголовок", 34), ("Ссылка", 10), ("Описание", 70),
)


def write_workbook(path: Path, offers: list[Offer], stats: dict[str, GroupStats]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    market = workbook.active
    market.title = "Рынок"
    market.append(["Группа", "Объявлений", "В расчёте", "Мин", "25%", "Медиана", "75%", "Макс"])
    for row in sorted(stats.values(), key=lambda s: (-s.count, s.group)):
        market.append([row.group, row.count, row.used, row.low, row.q1, row.median, row.q3, row.high])
    for index, width in enumerate((26, 11, 10, 10, 10, 11, 10, 10), start=1):
        market.column_dimensions[get_column_letter(index)].width = width
    for cell in market[1]:
        cell.font = Font(bold=True)

    sheet = workbook.create_sheet("Объявления")
    sheet.append([name for name, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    link_column = [name for name, _ in COLUMNS].index("Ссылка") + 1
    ordered = sorted(
        offers,
        key=lambda o: (0 if o.listing.is_new else 1, o.group, o.listing.price or 0),
    )
    for offer in ordered:
        listing = offer.listing
        sheet.append([
            "новое" if listing.is_new else "",
            offer.group,
            listing.price,
            offer.market,
            f"{offer.vs_market:+.0%}" if offer.vs_market is not None else "",
            offer.avito_price_badge,
            "; ".join(offer.notes),
            offer.seller,
            listing.closed_listings,
            listing.published,
            apartments_short_address(listing),
            listing.title,
            listing.url,
            listing.description,
        ])
        link = sheet.cell(row=sheet.max_row, column=link_column)
        link.hyperlink = listing.url
        link.value = "открыть"
        link.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def apartments_short_address(listing: apartments.Listing) -> str:
    return re.sub(r"^[^,]+,\s*", "", listing.address) if "," in listing.address else listing.address


def price_text(value: int | None) -> str:
    return f"{value:,} ₽".replace(",", " ") if value else "—"


def alert_fields(offer: Offer) -> dict[str, Any]:
    listing = offer.listing
    return {
        "id": listing.id,
        "title": listing.title,
        "price": price_text(listing.price),
        "market": price_text(offer.market),
        "vs_market": f"{offer.vs_market:+.0%}" if offer.vs_market is not None else "",
        "avito_badge": offer.avito_price_badge,
        "seller": offer.seller,
        "closed_listings": listing.closed_listings,
        "notes": offer.notes,
        "published": listing.published,
        "address": apartments_short_address(listing),
        "url": listing.url,
    }
