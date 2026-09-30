#!/usr/bin/env python3
"""Collect Saint Petersburg 2-room flats sold by owners into an Excel file.

The catalog search itself (city, rooms, "Частные" sellers) and all Avito
protection handling live in ``main.py``. This module turns the collected
listings into call-ready rows, separates agents from owners and remembers
which listings were already seen, so every run marks the new ones.

The first run walks the whole catalog. If Avito stops it half way, the next
run continues from the page where it stopped. Once the catalog has been
walked completely, runs only read pages until one holds nothing new.

    uv run python apartments.py          # continue / look for new listings
    uv run python apartments.py --full   # walk the whole catalog again
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from curl_cffi.requests.exceptions import (
    ConnectionError as CurlConnectionError,
    IncompleteRead,
    Timeout as CurlTimeout,
)

import main

LOGGER = logging.getLogger(__name__)
BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
SEEN_PATH = BASE_DIR / "data" / "seen_listings.json"
STATE_PATH = BASE_DIR / "data" / "scan_state.json"
# Raw catalog items of every listing found so far, keyed by listing id. The
# table is rebuilt from it, so improved rules apply to old listings as well.
LISTINGS_PATH = BASE_DIR / "data" / "listings.json"

VERDICT_OWNER = "собственник"
VERDICT_CHECK = "проверить"
VERDICT_AGENT = "агент"

# Avito's own seller-type line under the listing ("Агентство", "Компания"...).
AGENT_SELLER_LABELS = ("агентство", "компания", "застройщик")

# A seller who closed this many listings on Avito (in any category) is more
# likely to be a professional than a private owner, though not certainly.
MANY_CLOSED_LISTINGS = 100

# Sentences with these words ask agents to stay away. They are owner signals,
# and their agent vocabulary must not count against the listing.
NO_AGENTS_SENTENCE = re.compile(
    r"не\s+беспоко|не\s+звонит|не\s+звоните|не\s+предл[ао]г|не\s+нужн|"
    r"не\s+интерес|не\s+рассматрива|не\s+нуждаюсь|не\s+обрывать|отключен|"
    r"без\s+(?:агент|риелтор|риэлтор|посредник|реальн|клиент|покупател)|"
    r"просьб|прошу|не\s+агентств|не\s+тратьте|при\s+наличии|"
    r"(?:реальн|действующ)\w*\s+(?:покупател|клиент)|без\s+покупател"
)
# The owner already works with their own agent: still an owner's listing,
# but a weaker lead for an agent's call.
HAS_OWN_AGENT = re.compile(
    r"(?:есть|у\s+меня)\s+(?:уже\s+)?(?:сво[йи]\s+)?(?:агент|риелтор|риэлтор)|"
    r"(?:агент|риелтор|риэлтор)\w*(?:\s+[\w-]+){0,4}\s+(?:уже\s+)?есть|"
    r"сопровожда\w*\s+(?:мой\s+|наш\s+)?(?:агент|риелтор|риэлтор)"
)
AGENT_WORD = re.compile(r"агент|риелтор|риэлтор|посредник|маклер")
# The owner offers a fee to an agent who brings a buyer: the best lead.
PAYS_AGENT = re.compile(
    r"(?:готов\w*|могу)\s+(?:\S+\s+){0,2}(?:заплатить|оплатить|выплатить|отблагодарить)|"
    r"(?:заплачу|оплачу|выплачу)\s+(?:\S+\s+){0,3}(?:комисси|вознагражд|агент|риелтор|риэлтор)|"
    r"вознагражд\w*\s+(?:\S+\s+){0,3}(?:агент|риелтор|риэлтор)"
)
# A negated sentence ("комиссию не плачу", "эксклюзивный договор не
# подписываю") describes what the owner refuses, not who the seller is.
NEGATED_SENTENCE = re.compile(r"(?<!\w)(?:не|нет|без|ни)(?!\w)|отсутств|исключа")

# Phrases an agency writes and an owner practically never does.
AGENT_PHRASES = (
    (re.compile(r"агентств\w*\s+недвижимост"), "«агентство недвижимости»"),
    (re.compile(r"наш\w*\s+(?:агентств|компани|офис|специалист|эксперт)"), "«наше агентство/компания/специалист»"),
    (re.compile(r"эксклюзивн\w*\s+(?:прав|договор\w*\s+с\s+(?:нами|нашим|агентств))"), "«эксклюзивные права/договор»"),
    (re.compile(r"комисси\w*\s+(?:нашего\s+)?агентства\s+(?:составляет\s+)?\d"), "указана комиссия агентства"),
    (re.compile(r"(?:наш\w*|бесплатн\w*)\s+(?:полн\w*\s+)?(?:юридическ\w*\s+)?сопровождени"), "«наше/бесплатное сопровождение сделки»"),
    (re.compile(r"бесплатн\w*\s+(?:консультац|подбор|юридическ|одобрени)"), "«бесплатная консультация/подбор»"),
    (re.compile(r"ипотечн\w*\s+брокер|одобрени\w*\s+ипотеки\s+за"), "ипотечный брокер"),
    (re.compile(r"trade[\s-]?in|трейд[\s-]?ин"), "«трейд-ин»"),
    (re.compile(r"(?:интересы|поручению)\s+собственник"), "представляет собственника"),
    (re.compile(r"(?:показ\w*|просмотр\w*)\s+(?:проводит|организ\w*)\s+(?:наш\w*\s+)?(?:специалист|агент|риелтор|риэлтор|менеджер)"), "показы проводит специалист"),
    (re.compile(r"(?:номер|код|id)\s+объекта|лот\s*№"), "номер объекта в базе"),
    (re.compile(r"подбер[её]м|поможем\s+(?:продать|купить|подобрать)"), "«подберём/поможем» от лица компании"),
)

OWNER_PHRASES = (
    (re.compile(r"собственник"), "пишет «собственник»"),
    (re.compile(r"без\s+посредник"), "«без посредников»"),
    (re.compile(r"продаю\s+(?:свою|собственную|нашу)"), "«продаю свою»"),
)

TITLE_PATTERN = re.compile(
    r"(?P<area>\d+(?:[.,]\d+)?)\s*м².*?(?P<floor>\d+)\s*/\s*(?P<floors>\d+)\s*эт"
)
CLOSED_LISTINGS_PATTERN = re.compile(r"(\d+)\s+заверш")


@dataclass
class Listing:
    """One flat listing reduced to the fields useful for a call."""

    id: int
    url: str
    title: str
    price: int | None
    area: float | None
    floor: str
    address: str
    metro: str
    published: str
    timestamp_ms: int | None
    seller_label: str
    closed_listings: int | None
    avito_badges: tuple[str, ...]
    calls_allowed: bool
    description: str
    verdict: str = VERDICT_OWNER
    reasons: list[str] = field(default_factory=list)
    asks_no_agents: bool = False
    has_own_agent: bool = False
    pays_agent: bool = False
    is_new: bool = False

    @property
    def price_per_meter(self) -> int | None:
        if not self.price or not self.area:
            return None
        return round(self.price / self.area)


def _iva_payloads(item: dict[str, Any], step: str) -> list[dict[str, Any]]:
    steps = (item.get("iva") or {}).get(step) or []
    return [
        entry.get("payload") or {}
        for entry in steps
        if isinstance(entry, dict)
    ]


def _listing_url(url_path: str) -> str:
    """Absolute listing URL without Avito's per-search ``context`` query."""
    return main.BASE_URL + urlsplit(url_path or "").path


def listing_from_item(item: dict[str, Any]) -> Listing:
    """Extract the call-relevant fields from one catalog item."""
    title = str(item.get("title") or "")
    title_match = TITLE_PATTERN.search(title)
    area = None
    floor = ""
    if title_match:
        area = float(title_match["area"].replace(",", "."))
        floor = f"{title_match['floor']}/{title_match['floors']}"

    price_value = (item.get("priceDetailed") or {}).get("value")
    price = int(price_value) if isinstance(price_value, (int, float)) else None

    geo = item.get("geo") or {}
    coords = item.get("coords") or {}
    address = str(coords.get("address_user") or geo.get("formattedAddress") or "")
    metro = ""
    references = geo.get("geoReferences") or []
    if references and isinstance(references[0], dict):
        metro = f"{references[0].get('content', '')}{references[0].get('after', '')}".strip()

    published = ""
    for payload in _iva_payloads(item, "DateInfoStep"):
        relative = payload.get("relative") or payload.get("absolute")
        if relative:
            published = str(relative)
            break

    seller_label = ", ".join(
        str(payload["value"])
        for payload in _iva_payloads(item, "SecondLineStep")
        if payload.get("value")
    )

    closed_listings = None
    closed_text = str(item.get("closedItemsText") or "")
    closed_match = CLOSED_LISTINGS_PATTERN.search(closed_text)
    if closed_match:
        closed_listings = int(closed_match[1])
    elif closed_text.startswith("Нет"):
        closed_listings = 0

    badges = tuple(
        str(badge.get("title"))
        for payload in _iva_payloads(item, "BadgeBarStep")
        for badge in payload.get("badges") or []
        if isinstance(badge, dict) and badge.get("title")
    )

    timestamp = item.get("sortTimeStamp")
    return Listing(
        id=int(item["id"]),
        url=_listing_url(str(item.get("urlPath") or "")),
        title=title,
        price=price,
        area=area,
        floor=floor,
        address=address,
        metro=metro,
        published=published,
        timestamp_ms=int(timestamp) if isinstance(timestamp, (int, float)) else None,
        seller_label=seller_label,
        closed_listings=closed_listings,
        avito_badges=badges,
        calls_allowed=bool((item.get("contacts") or {}).get("phone")),
        description=str(item.get("description") or ""),
    )


def _sentences(text: str) -> list[str]:
    return [part for part in re.split(r"[.!?\n]+", text) if part.strip()]


def classify(listing: Listing) -> Listing:
    """Decide whether the listing is an owner's, an agent's, or unclear.

    Avito's "Собственник / Проверено в Росреестре" badge is deliberately
    ignored: it confirms the flat's owner and is shown on agency listings too.
    """
    agent_reasons: list[str] = []
    check_reasons: list[str] = []
    owner_reasons: list[str] = []

    label = listing.seller_label.lower()
    for agent_label in AGENT_SELLER_LABELS:
        if agent_label in label:
            agent_reasons.append(f"Авито помечает продавца: «{listing.seller_label}»")
            break

    text = listing.description.lower().replace("ё", "е")
    scan_parts = []
    for sentence in _sentences(text):
        if HAS_OWN_AGENT.search(sentence):
            listing.has_own_agent = True
            continue
        if AGENT_WORD.search(sentence) and NO_AGENTS_SENTENCE.search(sentence):
            listing.asks_no_agents = True
            continue
        scan_parts.append(sentence)
    scan_text = ". ".join(scan_parts)

    listing.pays_agent = bool(PAYS_AGENT.search(text))
    positive_text = ". ".join(
        sentence for sentence in scan_parts if not NEGATED_SENTENCE.search(sentence)
    )
    for pattern, reason in AGENT_PHRASES:
        if pattern.search(positive_text):
            agent_reasons.append(f"в описании: {reason}")
    # Third-person talk about "the owners" who do not work with other agents
    # is how an exclusive agent writes, even though the sentence is negated.
    if re.search(r"собственник\w*\s+(?:с\s+)?други\w*\s+агент\w*\s+не\s+сотруднича", scan_text):
        agent_reasons.append("в описании: «собственник с другими агентами не сотрудничает»")

    if (
        listing.closed_listings is not None
        and listing.closed_listings >= MANY_CLOSED_LISTINGS
    ):
        check_reasons.append(
            f"у продавца {listing.closed_listings} завершённых объявлений"
        )

    if listing.pays_agent:
        owner_reasons.append("готов платить агенту за покупателя")
    if listing.has_own_agent:
        owner_reasons.append("у собственника уже есть свой агент")
    if listing.asks_no_agents:
        owner_reasons.append("просит агентов не беспокоить")
    for pattern, reason in OWNER_PHRASES:
        if pattern.search(scan_text):
            owner_reasons.append(reason)

    if agent_reasons:
        listing.verdict = VERDICT_AGENT
        listing.reasons = agent_reasons + check_reasons
    elif check_reasons:
        listing.verdict = VERDICT_CHECK
        listing.reasons = check_reasons + owner_reasons
    else:
        listing.verdict = VERDICT_OWNER
        listing.reasons = owner_reasons or ["признаков агента нет"]
    return listing


def load_seen(path: Path) -> dict[str, str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}
    except (OSError, json.JSONDecodeError) as exc:
        LOGGER.warning("cannot read %s (%s); treating all listings as new", path, exc)
        return {}
    return data if isinstance(data, dict) else {}


def mark_new(listings: list[Listing], seen: dict[str, str], now: datetime) -> dict[str, str]:
    """Flag listings absent from ``seen`` and return the updated registry."""
    updated = dict(seen)
    stamp = now.isoformat(timespec="seconds")
    for listing in listings:
        key = str(listing.id)
        listing.is_new = key not in seen
        updated.setdefault(key, stamp)
    return updated


def save_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=0), encoding="utf-8")
    tmp.replace(path)


def load_state(path: Path) -> dict[str, Any]:
    """Progress of the full catalog walk: ``{"complete": bool, "next_page": int}``."""
    state = load_seen(path)
    next_page = state.get("next_page")
    return {
        "complete": state.get("complete") is True,
        "next_page": next_page if isinstance(next_page, int) and next_page >= 1 else 1,
    }


@dataclass
class ScanOutcome:
    items: list[dict[str, Any]]
    error: str | None
    last_ok_page: int | None
    reached_end: bool


COLUMNS = (
    ("Новое", 8),
    ("Вердикт", 13),
    ("Почему", 45),
    ("Агентам не звонить", 11),
    ("Свой агент уже есть", 11),
    ("Готов платить агенту", 11),
    ("Звонки", 10),
    ("Опубликовано", 14),
    ("Цена, ₽", 13),
    ("Площадь, м²", 9),
    ("₽ за м²", 10),
    ("Этаж", 7),
    ("Адрес", 40),
    ("Метро", 22),
    ("Пометка Авито", 13),
    ("Завершённых объявлений", 11),
    ("Значки Авито", 30),
    ("Заголовок", 30),
    ("Ссылка", 16),
    ("Описание", 80),
)


def _row(listing: Listing) -> list[Any]:
    return [
        "новое" if listing.is_new else "",
        listing.verdict,
        "; ".join(listing.reasons),
        "да" if listing.asks_no_agents else "",
        "да" if listing.has_own_agent else "",
        "да" if listing.pays_agent else "",
        "да" if listing.calls_allowed else "только чат",
        listing.published,
        listing.price,
        listing.area,
        listing.price_per_meter,
        listing.floor,
        listing.address,
        listing.metro,
        listing.seller_label,
        listing.closed_listings,
        ", ".join(listing.avito_badges),
        listing.title,
        listing.url,
        listing.description,
    ]


def _sort_key(listing: Listing) -> tuple[int, int, int]:
    verdict_rank = {VERDICT_OWNER: 0, VERDICT_CHECK: 1, VERDICT_AGENT: 2}
    return (
        0 if listing.is_new else 1,
        verdict_rank.get(listing.verdict, 3),
        -(listing.timestamp_ms or 0),
    )


def write_workbook(path: Path, listings: list[Listing]) -> None:
    """Owners and unclear cases on the first sheet, rejected agents on the second."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    ordered = sorted(listings, key=_sort_key)
    sheets = (
        ("Собственники", [x for x in ordered if x.verdict != VERDICT_AGENT]),
        ("Отсеяно — агенты", [x for x in ordered if x.verdict == VERDICT_AGENT]),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    link_column = [name for name, _ in COLUMNS].index("Ссылка") + 1
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append([name for name, _ in COLUMNS])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index, (_, width) in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for listing in rows:
            sheet.append(_row(listing))
            link = sheet.cell(row=sheet.max_row, column=link_column)
            link.hyperlink = listing.url
            link.value = "открыть"
            link.font = Font(color="0563C1", underline="single")
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def page_has_nothing_new(result: main.PageRequestResult, seen: dict[str, str]) -> bool:
    ids = [item.get("id") for item in result.items]
    return bool(ids) and all(str(item_id) in seen for item_id in ids)


def collect_listings(
    *,
    start_page: int = 1,
    seen: dict[str, str] | None = None,
) -> ScanOutcome:
    """Run the catalog search from ``start_page``.

    With ``seen`` the walk ends at the first page without unseen listings.
    """
    items: dict[int, dict[str, Any]] = {}
    ok_pages: list[int] = []
    reached_end = False

    def keep(results: tuple[main.PageRequestResult, ...]) -> None:
        nonlocal reached_end
        for result in results:
            if result.status_code != 200:
                continue
            ok_pages.append(result.page)
            if result.stats is not None and main.is_last_catalog_page(
                result.page, result.stats
            ):
                reached_end = True
            for item in result.items:
                if isinstance(item.get("id"), int):
                    items[item["id"]] = item
        LOGGER.info("collected %s unique listings so far", len(items))

    stop_after = None
    if seen is not None:
        def stop_after(result: main.PageRequestResult) -> bool:
            return page_has_nothing_new(result, seen)

    error = None
    try:
        main.run(on_page_results=keep, start_page=start_page, stop_after=stop_after)
    except (RuntimeError, IncompleteRead, CurlConnectionError, CurlTimeout) as exc:
        error = str(exc)
        LOGGER.error("search stopped early: %s", exc)
    return ScanOutcome(
        items=list(items.values()),
        error=error,
        last_ok_page=max(ok_pages) if ok_pages else None,
        reached_end=reached_end,
    )


def configure_logging() -> None:
    console = logging.StreamHandler()
    console.setLevel(logging.INFO)
    file_handler = logging.FileHandler("firewall-debug.log", encoding="utf-8")
    file_handler.setLevel(logging.DEBUG)
    logging.basicConfig(
        level=logging.DEBUG,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[console, file_handler],
    )


@dataclass
class TableResult:
    path: Path
    listings: list[Listing]


def build_table(base: dict[str, Any], seen: dict[str, str], now: datetime) -> TableResult:
    """Classify every stored listing, mark the new ones and write the workbook."""
    listings = [classify(listing_from_item(item)) for item in base.values()]
    save_json(SEEN_PATH, mark_new(listings, seen, now))
    path = OUTPUT_DIR / f"двушки-спб-собственники-{now:%Y-%m-%d_%H-%M}.xlsx"
    write_workbook(path, listings)
    return TableResult(path=path, listings=listings)


def next_state(
    state: dict[str, Any], outcome: ScanOutcome, *, full_walk: bool
) -> dict[str, Any]:
    """Where the next run should continue."""
    if not full_walk:
        return state
    if outcome.reached_end and outcome.error is None:
        return {"complete": True, "next_page": 1}
    if outcome.last_ok_page is not None:
        return {"complete": False, "next_page": outcome.last_ok_page + 1}
    return state


def run_search(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--full",
        action="store_true",
        help="пройти весь каталог заново с первой страницы",
    )
    args = parser.parse_args(argv)
    configure_logging()

    seen = load_seen(SEEN_PATH)
    state = load_state(STATE_PATH)
    if args.full:
        state = {"complete": False, "next_page": 1}
    full_walk = not state["complete"]
    if full_walk:
        print(f"Полный обход каталога, начиная со страницы {state['next_page']}.")
        outcome = collect_listings(start_page=state["next_page"])
    else:
        print("Каталог уже пройден целиком; ищу только новые объявления.")
        outcome = collect_listings(seen=seen)
    save_json(STATE_PATH, next_state(state, outcome, full_walk=full_walk))

    error = outcome.error
    base = load_seen(LISTINGS_PATH)
    for item in outcome.items:
        base[str(item["id"])] = item
    if not base:
        print("Не удалось получить ни одного объявления. Подробности в firewall-debug.log.")
        return 1
    save_json(LISTINGS_PATH, base)

    now = datetime.now()
    output = build_table(base, seen, now)
    listings = output.listings

    counts = {
        verdict: sum(listing.verdict == verdict for listing in listings)
        for verdict in (VERDICT_OWNER, VERDICT_CHECK, VERDICT_AGENT)
    }
    new_owners = sum(
        listing.is_new and listing.verdict != VERDICT_AGENT for listing in listings
    )
    print()
    if error:
        print(f"Авито остановил поиск: {error}")
        print("Сохранено то, что успели собрать.")
        if full_walk and outcome.last_ok_page is not None:
            print(
                "Следующий запуск продолжит со страницы "
                f"{outcome.last_ok_page + 1}; подождите 15–30 минут."
            )
    print(f"Найдено за этот запуск: {len(outcome.items)}")
    print(f"Всего в базе: {len(listings)}")
    print(f"  собственники: {counts[VERDICT_OWNER]}")
    print(f"  проверить вручную: {counts[VERDICT_CHECK]}")
    print(f"  агенты (отсеяны): {counts[VERDICT_AGENT]}")
    print(f"Новых (не агентов) с прошлого запуска: {new_owners}")
    print(f"Таблица: {output.path}")
    return 0


if __name__ == "__main__":
    sys.exit(run_search())
